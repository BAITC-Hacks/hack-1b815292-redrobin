"""Tests for DOCX parsing and clause segmentation."""

import hashlib
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document as OpenDocument

from app.models.domain import Document, DocumentRole
from app.services.docx_parser import DocxParser

SOURCE_DIR = Path("materials/source")
BEFORE_PATH = SOURCE_DIR / "Положение_о_внутреннем_аудите_редакция_8_обезличено.docx"
AFTER_PATH = SOURCE_DIR / "Положение_о_внутреннем_аудите_редакция_9_обезличено.docx"


def source_document(path: Path, role: DocumentRole) -> Document:
    return Document(
        id=uuid4(),
        role=role,
        original_name=path.name,
        safe_path=str(path),
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        size_bytes=path.stat().st_size,
    )


@pytest.mark.parametrize(
    ("path", "role"),
    [
        (BEFORE_PATH, DocumentRole.BEFORE),
        (AFTER_PATH, DocumentRole.AFTER),
    ],
)
def test_real_documents_include_priority_clauses(
    path: Path, role: DocumentRole
) -> None:
    parsed = DocxParser().parse(source_document(path, role))
    clauses_by_number = {
        clause.number: clause for clause in parsed.clauses if clause.number
    }

    for number in ("3.4", "3.9", "5.3", "9.15"):
        assert number in clauses_by_number
        clause = clauses_by_number[number]
        assert clause.raw_text.startswith(f"{number}.")
        assert clause.source_end > clause.source_start
    assert parsed.document.paragraph_count > 0
    assert parsed.sections


def test_glued_clauses_are_split_with_stable_ranges() -> None:
    parsed = DocxParser().parse(source_document(BEFORE_PATH, DocumentRole.BEFORE))
    clauses = {
        clause.number: clause
        for clause in parsed.clauses
        if clause.number in {"3.9", "3.10", "3.11"}
    }

    assert set(clauses) == {"3.9", "3.10", "3.11"}
    assert len({clause.id for clause in clauses.values()}) == 3
    assert len({clause.paragraph_index for clause in clauses.values()}) == 1
    assert clauses["3.9"].source_end <= clauses["3.10"].source_start
    assert clauses["3.10"].source_end <= clauses["3.11"].source_start


def test_table_rows_are_not_lost(tmp_path: Path) -> None:
    source_path = tmp_path / "table.docx"
    source = OpenDocument()
    source.add_heading("1. Проверка", level=1)
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Роль"
    table.cell(0, 1).text = "Функция"
    table.cell(1, 0).text = "Аудитор"
    table.cell(1, 1).text = "Проверяет контроль"
    source.save(source_path)

    parsed = DocxParser().parse(source_document(source_path, DocumentRole.BEFORE))
    table_clauses = [
        clause for clause in parsed.clauses if clause.style_name == "TableRow"
    ]

    assert len(table_clauses) == 2
    assert "Роль | Функция" in table_clauses[0].raw_text
    assert "Аудитор | Проверяет контроль" in table_clauses[1].raw_text
