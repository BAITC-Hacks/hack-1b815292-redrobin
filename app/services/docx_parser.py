"""Deterministic DOCX and OOXML parsing."""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from docx import Document as OpenDocument
from docx.document import Document as OpenDocumentType
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.errors import AppError
from app.models.domain import Document, DocumentStatus, ParsedDocument
from app.services.clause_segmenter import ClauseSegmenter, clause_segmenter


@dataclass(frozen=True)
class DocumentBlock:
    """An ordered paragraph or flattened table row from the source DOCX."""

    index: int
    text: str
    style_name: str | None
    numbering_metadata: dict[str, str] | None
    kind: str


class DocxParser:
    """Read DOCX blocks without invoking AI or following external content."""

    def __init__(self, segmenter: ClauseSegmenter | None = None) -> None:
        self.segmenter = segmenter or clause_segmenter

    def parse(self, document: Document) -> ParsedDocument:
        try:
            source = OpenDocument(Path(document.safe_path))
            blocks = list(self._extract_blocks(source))
        except (PackageNotFoundError, KeyError, ValueError, OSError) as exc:
            document.status = DocumentStatus.FAILED
            raise AppError(
                422,
                "invalid_docx",
                "Файл поврежден или не является DOCX.",
                field=f"{document.role.value}_file",
            ) from exc

        sections, clauses, warnings = self.segmenter.segment(document.id, blocks)
        document.paragraph_count = len(blocks)
        document.status = DocumentStatus.PARSED
        return ParsedDocument(
            document=document,
            sections=sections,
            clauses=clauses,
            warnings=warnings,
        )

    def _extract_blocks(self, document: OpenDocumentType) -> Iterator[DocumentBlock]:
        block_index = 0
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, document)
                text = paragraph.text
                if text and text.strip():
                    yield DocumentBlock(
                        index=block_index,
                        text=text,
                        style_name=self._style_name(paragraph),
                        numbering_metadata=self._numbering_metadata(paragraph),
                        kind="paragraph",
                    )
                    block_index += 1
            elif isinstance(child, CT_Tbl):
                table = Table(child, document)
                for row in table.rows:
                    cells = [self._cell_text(cell.paragraphs) for cell in row.cells]
                    text = " | ".join(cell for cell in cells if cell)
                    if text.strip():
                        yield DocumentBlock(
                            index=block_index,
                            text=text,
                            style_name="TableRow",
                            numbering_metadata={"source": "table"},
                            kind="table_row",
                        )
                        block_index += 1

    @staticmethod
    def _cell_text(paragraphs: list[Paragraph]) -> str:
        return " ".join(
            paragraph.text.strip()
            for paragraph in paragraphs
            if paragraph.text and paragraph.text.strip()
        )

    @staticmethod
    def _style_name(paragraph: Paragraph) -> str | None:
        try:
            return paragraph.style.name if paragraph.style is not None else None
        except (AttributeError, KeyError):
            return None

    @staticmethod
    def _numbering_metadata(paragraph: Paragraph) -> dict[str, str] | None:
        properties = paragraph._p.pPr
        if properties is None or properties.numPr is None:
            return None
        metadata: dict[str, str] = {}
        if properties.numPr.numId is not None:
            metadata["num_id"] = str(properties.numPr.numId.val)
        if properties.numPr.ilvl is not None:
            metadata["level"] = str(properties.numPr.ilvl.val)
        return metadata or None


docx_parser = DocxParser()
